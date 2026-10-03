"""PDF, Word, and PowerPoint readers.

Each needs an optional dependency, so each reports its own availability with an
actionable install hint rather than the file being silently skipped. Locators
are the natural address of the format — ``p.4``, a heading path, ``slide 7`` —
so a citation lands somewhere a human can actually look.

A PDF with no extractable text is a scanned image, not an empty document. That
is reported as such (OCR required) instead of contributing zero evidence while
appearing to have been read.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from automation_miner.readers.base import (
    Availability,
    BaseReader,
    MediaType,
    MetaValue,
    ReaderError,
    Segment,
    missing_dependency,
)

MAX_TABLE_ROWS = 50
PDF_SUSPICIOUS_TOKEN_MIN = 20
PDF_SUSPICIOUS_TOKEN_RATIO = 0.02


def pdf_extraction_quality(text: str) -> tuple[float, list[str]]:
    """Estimate whether a PDF's character map corrupted otherwise visible prose.

    Broken ToUnicode maps commonly turn words into ``operaDng`` or ``por;olio``.
    This is not mojibake that an encoding retry can repair. A conservative
    threshold avoids flagging a handful of camelCase identifiers in technical
    PDFs while catching document-wide substitution patterns.
    """
    tokens = re.findall(r"\S+", text)
    if not tokens:
        return 0.0, []
    suspicious = [
        token
        for token in tokens
        if re.search(r"[a-z][A-Z]", token)
        or re.search(r"[A-Za-z][;\"|][A-Za-z]", token)
        or "\ufffd" in token
    ]
    ratio = len(suspicious) / len(tokens)
    quality = max(0.0, round(1.0 - ratio, 3))
    return quality, suspicious[:8]


_SUSPICIOUS = (
    re.compile(r"[a-z][A-Z]"),
    re.compile(r"[A-Za-z][;\"|][A-Za-z]"),
)


def _suspicious_count(segments: list[Segment]) -> int:
    return sum(
        1
        for segment in segments
        for token in re.findall(r"\S+", segment.text)
        if any(p.search(token) for p in _SUSPICIOUS) or "\ufffd" in token
    )


def _open_pdf(data: bytes) -> object:
    import pypdf

    try:
        document = pypdf.PdfReader(io.BytesIO(data))
    except Exception as exc:  # pypdf raises several unrelated error types
        raise ReaderError(f"cannot open PDF: {exc}") from exc
    if document.is_encrypted:
        try:
            opened = bool(document.decrypt(""))
        except Exception:
            opened = False
        if not opened:
            raise ReaderError("PDF is password-protected")
    return document


def _pdf_pages(pages: list[object]) -> tuple[list[Segment], list[str]]:
    """One segment per page with text; every failed or empty page recorded."""
    segments: list[Segment] = []
    errors: list[str] = []
    for number, page in enumerate(pages, 1):
        try:
            text = getattr(page, "extract_text")() or ""
        except Exception as exc:
            errors.append(f"p.{number}: {type(exc).__name__}: {exc}")
            continue
        if text.strip():
            segments.append(Segment(text=text, locator=f"p.{number}"))
        else:
            errors.append(f"p.{number}: no extractable text")
    return segments, errors


class PdfReader(BaseReader):
    """PDF text extraction via pypdf, one segment per page."""

    name = "pdf"
    suffixes = (".pdf",)
    media_type = MediaType.DOCUMENT
    binary = True

    def available(self) -> Availability:
        return missing_dependency("pypdf", "pdf")

    def parse(self, path: Path, data: bytes) -> list[Segment]:
        pages = list(getattr(_open_pdf(data), "pages"))
        segments, errors = _pdf_pages(pages)
        self.note("pages", len(pages))
        self.note("pages_failed", len(errors))
        if errors:
            self.note("page_errors", "; ".join(errors[:8]))
        if not segments:
            raise ReaderError(
                f"no extractable text in {len(pages)} page(s) — likely a scanned "
                "PDF; OCR it before mining",
                meta={
                    "pages": len(pages), "pages_with_text": 0, "pages_failed": len(errors),
                    "page_errors": "; ".join(errors[:8]),
                },
            )
        self._check_quality(segments)
        return segments

    def _check_quality(self, segments: list[Segment]) -> None:
        quality, examples = pdf_extraction_quality("\n".join(s.text for s in segments))
        self.note("extraction_quality", quality)
        corrupted = (
            examples
            and (1.0 - quality) >= PDF_SUSPICIOUS_TOKEN_RATIO
            and _suspicious_count(segments) >= PDF_SUSPICIOUS_TOKEN_MIN
        )
        if corrupted and not bool(self.options.get("allow_low_quality", False)):
            raise ReaderError(
                "low-quality PDF text extraction detected "
                f"(quality {quality:.1%}; examples: {', '.join(examples)}). OCR the PDF "
                "or set [readers.pdf] allow_low_quality = true to override"
            )

    def meta(self, path: Path, data: bytes, segments: list[Segment]) -> dict[str, MetaValue]:
        return {"pages_with_text": len(segments), **super().meta(path, data, segments)}


def _docx_sections(paragraphs: list[object]) -> list[Segment]:
    """Paragraphs grouped under their heading path."""
    segments: list[Segment] = []
    stack: list[tuple[int, str]] = []
    buffer: list[str] = []
    locator = ""
    for paragraph in paragraphs:
        text = str(getattr(paragraph, "text", "")).strip()
        level = DocxReader.heading_level(paragraph)
        if level and text:
            if buffer:
                segments.append(Segment(text="\n".join(buffer), locator=locator))
                buffer = []
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, text))
            locator = " > ".join(title for _, title in stack)
        if text:
            buffer.append(text)
    if buffer:
        segments.append(Segment(text="\n".join(buffer), locator=locator))
    return segments


def _docx_table(table: object, index: int) -> Segment | None:
    all_rows = list(getattr(table, "rows"))
    rows = [
        " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells)
        for row in all_rows[:MAX_TABLE_ROWS]
    ]
    rows = [row for row in rows if row.replace("|", "").strip()]
    if not rows:
        return None
    extra = len(all_rows) - MAX_TABLE_ROWS
    note = f"\n(+{extra} further rows)" if extra > 0 else ""
    return Segment(text="\n".join(rows) + note, locator=f"table {index}")


class DocxReader(BaseReader):
    """Word documents via python-docx, segmented on heading styles."""

    name = "docx"
    suffixes = (".docx",)
    media_type = MediaType.DOCUMENT
    binary = True

    def available(self) -> Availability:
        return missing_dependency("docx", "docx")

    def parse(self, path: Path, data: bytes) -> list[Segment]:
        import docx

        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ReaderError(f"cannot open document: {exc}") from exc
        segments = _docx_sections(list(document.paragraphs))
        self.note("tables", len(document.tables))
        tables = (_docx_table(t, i) for i, t in enumerate(document.tables, 1))
        segments += [segment for segment in tables if segment is not None]
        if not segments:
            raise ReaderError("document contains no extractable text")
        return segments

    @staticmethod
    def heading_level(paragraph: object) -> int:
        name = getattr(getattr(paragraph, "style", None), "name", "") or ""
        if not name.lower().startswith("heading"):
            return 0
        tail = name.split()[-1]
        return int(tail) if tail.isdigit() else 1


def _slide_text(slide: object) -> list[str]:
    """Text frames, table rows, and speaker notes of one slide."""
    parts: list[str] = []
    for shape in getattr(slide, "shapes"):
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                cells = " | ".join(c.text.strip() for c in row.cells)
                if cells.replace("|", "").strip():
                    parts.append(cells)
    if getattr(slide, "has_notes_slide"):
        frame = getattr(getattr(slide, "notes_slide"), "notes_text_frame")
        notes = str(getattr(frame, "text", "")).strip()
        if notes:
            parts.append(f"Speaker notes: {notes}")
    return parts


class PptxReader(BaseReader):
    """PowerPoint decks via python-pptx, one segment per slide plus notes."""

    name = "pptx"
    suffixes = (".pptx",)
    media_type = MediaType.PRESENTATION
    binary = True

    def available(self) -> Availability:
        return missing_dependency("pptx", "pptx")

    def parse(self, path: Path, data: bytes) -> list[Segment]:
        from pptx import Presentation

        try:
            deck = Presentation(io.BytesIO(data))
        except Exception as exc:
            raise ReaderError(f"cannot open presentation: {exc}") from exc
        slides = list(deck.slides)
        self.note("slides", len(slides))
        segments = [
            Segment(text="\n".join(parts), locator=f"slide {number}")
            for number, slide in enumerate(slides, 1)
            if (parts := _slide_text(slide))
        ]
        if not segments:
            raise ReaderError(f"no extractable text in {len(slides)} slide(s)")
        return segments
